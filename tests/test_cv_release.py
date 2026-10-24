"""Regression tests for image-only PDFs and non-image CV routing."""
import io
from unittest.mock import AsyncMock, patch

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter

from backend.app import app


def blank_pdf():
    data = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.write(data)
    return data.getvalue()


def test_image_only_pdf_requires_pdf_capable_provider():
    client = TestClient(app)
    with patch('backend.routes.cv_upload.claude_configured', return_value=False), \
         patch('backend.routes.cv_upload.extract_suggestions', new_callable=AsyncMock) as extract:
        response = client.post(
            '/api/extract',
            files={'file': ('synthetic_scan.pdf', blank_pdf(), 'application/pdf')},
            data={'provider': 'auto'},
        )

    assert response.status_code == 422
    assert 'no readable text' in response.json()['detail'].lower()
    extract.assert_not_awaited()


def test_image_only_pdf_uses_claude_in_auto_mode():
    client = TestClient(app)
    with patch('backend.routes.cv_upload.claude_configured', return_value=True), \
         patch(
             'backend.routes.cv_upload.extract_suggestions',
             new_callable=AsyncMock,
             return_value={'provider': 'claude', 'candidates': [], 'scanned': True,
                           'requires_confirmation': True, 'stored': False},
         ) as extract:
        response = client.post(
            '/api/extract',
            files={'file': ('synthetic_scan.pdf', blank_pdf(), 'application/pdf')},
            data={'provider': 'auto'},
        )

    assert response.status_code == 200
    assert response.json()['provider'] == 'claude'
    assert response.json()['scanned'] is True
    assert extract.await_args.args[1] == 'claude'


def test_explicit_runpod_rejects_image_only_pdf():
    client = TestClient(app)
    response = client.post(
        '/api/extract',
        files={'file': ('synthetic_scan.pdf', blank_pdf(), 'application/pdf')},
        data={'provider': 'runpod'},
    )
    assert response.status_code == 422


def test_text_docx_keeps_auto_provider_selection():
    doc = Document()
    doc.add_paragraph('I wrote SQL queries and Python scripts for reports.')
    data = io.BytesIO()
    doc.save(data)
    client = TestClient(app)

    with patch(
        'backend.routes.cv_upload.extract_suggestions',
        new_callable=AsyncMock,
        return_value={'provider': 'dictionary', 'candidates': [], 'scanned': False,
                      'requires_confirmation': True, 'stored': False},
    ) as extract:
        response = client.post(
            '/api/extract',
            files={'file': ('synthetic_cv.docx', data.getvalue(),
                            'application/vnd.openxmlformats-officedocument.wordprocessingml.document')},
            data={'provider': 'auto'},
        )

    assert response.status_code == 200
    assert 'Python scripts' in extract.await_args.args[0]
    assert extract.await_args.args[1] == 'auto'



def test_docx_table_skills_are_read_in_document_order():
    from backend.services.documents import parse_cv
    doc = Document()
    doc.add_paragraph('PROJECTS')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Technology'
    table.cell(0, 1).text = 'Experience'
    table.cell(1, 0).text = 'Laravel and PostgreSQL'
    table.cell(1, 1).text = 'Built a full-stack application'
    doc.add_paragraph('Used JavaScript for frontend logic.')
    output = io.BytesIO()
    doc.save(output)
    parsed = parse_cv('fictional.docx', output.getvalue())
    assert parsed.index('PROJECTS') < parsed.index('Laravel and PostgreSQL')
    assert parsed.index('Built a full-stack application') < parsed.index('JavaScript')
    assert 'Technology | Experience' in parsed
    assert '\n' in parsed


def test_empty_docx_rejected_without_gpu_inference():
    doc = Document()
    doc.add_paragraph('   ')
    output = io.BytesIO()
    doc.save(output)
    with patch('backend.routes.cv_upload.extract_suggestions', new_callable=AsyncMock) as ai:
        response = TestClient(app).post(
            '/api/extract',
            files={'file': ('fictional.docx', output.getvalue(),
                            'application/vnd.openxmlformats-officedocument.wordprocessingml.document')},
            data={'provider': 'auto'},
        )
    assert response.status_code == 422
    assert 'no readable text' in response.json()['detail'].lower()
    ai.assert_not_awaited()


def test_bad_pdf_rejected_without_optional_claude_billing():
    with patch('backend.routes.cv_upload.claude_configured', return_value=True), \
         patch('backend.routes.cv_upload.extract_suggestions', new_callable=AsyncMock) as ai:
        response = TestClient(app).post(
            '/api/extract',
            files={'file': ('broken.pdf', b'%PDF-1.7 invalid PDF data', 'application/pdf')},
            data={'provider': 'auto'},
        )
    assert response.status_code == 400
    assert 'could not read pdf' in response.json()['detail'].lower()
    ai.assert_not_awaited()
